#
# Copyright 2021 Open Raven Inc. and the Mockingbird project authors
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     https://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#
import random
from typing import final

from openpyxl import Workbook

from .__base import __BasePandaDocument
from ..__base import __BaseDocument, __BaseStructuredDataType


class XLSXDocument(__BaseDocument):
    EXT = "xlsx"

    def __init__(self, config_file=None):
        super().__init__(extension=XLSXDocument.EXT, config_file=config_file)

        # Create a list of formats we're going to save xlsx files in.
        self._xlsx_styles = []
        active_styles = self._configurable_dict["structured_data"]["xlsx_document"]["active_styles"]

        if active_styles["pandas_xlsx_writer"]:
            self._xlsx_styles.append(_XlsxDocumentPandasXlsxWriterStyle)

        if active_styles["openpyxl"]:
            self._xlsx_styles.append(_XlsxDocumentOpenPyxlStyle)

    @final
    def save(self, save_path: str) -> None:

        for style in self._xlsx_styles:
            instantiated_style = style(config_file=self._config_file)
            instantiated_style.clone_sensitive_data(other=self)
            instantiated_style.save(save_path=save_path)
            self._meta_data_object.add_other_meta_data(instantiated_style._meta_data_object)


class _XlsxDocumentOpenPyxlStyle(__BaseStructuredDataType):
    @final
    def __init__(self, config_file=None):
        super().__init__(extension=XLSXDocument.EXT, config_file=config_file)

    @final
    def save(self, save_path: str) -> None:
        """
        Writes the structured data into a xlsx file. Fills the other pages in the excel sheet with random jibberish, but
        one of the pages has a pii-entry hidden within it somewhere.
        """

        # How many pages for the excel file
        pages = random.randint(1, 10)  # todo, add to config / support for this
        pii_page = random.randint(0, pages - 1)

        save_file = self.setup_save_file(save_path=save_path, extension=self.extension)
        wb = Workbook()

        for x in range(pages):
            ws = wb.create_sheet("mysheet", x)

            first_row = True

            if x == pii_page:
                structured_array = self._get_structured_data()
            else:
                structured_array = self._get_structured_data_no_sensitive_info()

            for line in structured_array:
                if first_row:
                    # write header first
                    header = list(line.keys())
                    ws.append(header)
                    first_row = False

                # export all line's values
                ws.append(list(line.values()))

        wb.save(save_file)
        self._log_save(save_file)


class _XlsxDocumentPandasXlsxWriterStyle(__BasePandaDocument):
    """
    Writes an xlsx document using Pandas dataframes and the xlsxwriter library.
    """

    def __init__(self, config_file=None):
        super().__init__(XLSXDocument.EXT, config_file)

    @final
    def save(self, save_path: str) -> None:
        save_file = self.setup_save_file(save_path=save_path, extension=self.extension)

        dataframe = self._get_data_frame()
        dataframe.to_excel(save_file, engine="xlsxwriter")

        self._log_save(save_file)
